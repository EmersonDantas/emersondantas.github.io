"""
Gera o currículo DOCX otimizado para ATS (Gupy, LinkedIn, etc.)
Boas práticas ATS:
- Sem tabelas, colunas ou layouts complexos
- Headings padrão para seções
- Texto plano, sem imagens ou gráficos
- Fontes padrão (Calibri)
- Ordem cronológica reversa
- Keywords relevantes em texto corrido
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE

doc = Document()

# --- Estilos ---
style = doc.styles["Normal"]
font = style.font
font.name = "Calibri"
font.size = Pt(10.5)
font.color.rgb = RGBColor(0x33, 0x33, 0x33)
style.paragraph_format.space_after = Pt(2)
style.paragraph_format.space_before = Pt(0)

for level in range(1, 4):
    heading_style = doc.styles[f"Heading {level}"]
    heading_style.font.name = "Calibri"
    heading_style.font.color.rgb = RGBColor(0x0B, 0x1D, 0x32)
    heading_style.paragraph_format.space_before = Pt(12 if level == 2 else 6)
    heading_style.paragraph_format.space_after = Pt(4)
    if level == 1:
        heading_style.font.size = Pt(22)
    elif level == 2:
        heading_style.font.size = Pt(14)
    else:
        heading_style.font.size = Pt(11)


def add_separator():
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run("_" * 80)
    run.font.size = Pt(2)
    run.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)


def add_experience(role, company, period, context=None, items=None, techs=None):
    h = doc.add_heading(role, level=3)
    h.paragraph_format.space_before = Pt(10)
    p = doc.add_paragraph()
    run_company = p.add_run(company)
    run_company.bold = True
    run_company.font.size = Pt(10.5)
    p.add_run(f"  |  {period}").font.size = Pt(10)

    if context:
        ctx = doc.add_paragraph(context)
        ctx.paragraph_format.space_before = Pt(2)
        ctx.style.font.size = Pt(10.5)

    if items:
        for item in items:
            bullet = doc.add_paragraph(item, style="List Bullet")
            bullet.paragraph_format.space_after = Pt(1)
            bullet.paragraph_format.space_before = Pt(0)
            for run in bullet.runs:
                run.font.size = Pt(10)

    if techs:
        tp = doc.add_paragraph()
        tp.paragraph_format.space_before = Pt(2)
        label = tp.add_run("Tecnologias: ")
        label.bold = True
        label.font.size = Pt(10)
        tp.add_run(", ".join(techs)).font.size = Pt(10)


def add_project(name, role, period, description, items=None, techs=None):
    h = doc.add_heading(name, level=3)
    h.paragraph_format.space_before = Pt(10)
    p = doc.add_paragraph()
    run_role = p.add_run(role)
    run_role.bold = True
    run_role.font.size = Pt(10.5)
    p.add_run(f"  |  {period}").font.size = Pt(10)

    desc = doc.add_paragraph(description)
    desc.paragraph_format.space_before = Pt(2)

    if items:
        for item in items:
            bullet = doc.add_paragraph(item, style="List Bullet")
            bullet.paragraph_format.space_after = Pt(1)
            bullet.paragraph_format.space_before = Pt(0)
            for run in bullet.runs:
                run.font.size = Pt(10)

    if techs:
        tp = doc.add_paragraph()
        tp.paragraph_format.space_before = Pt(2)
        label = tp.add_run("Tecnologias: ")
        label.bold = True
        label.font.size = Pt(10)
        tp.add_run(", ".join(techs)).font.size = Pt(10)


# ==========================================
# CABEÇALHO
# ==========================================
title = doc.add_heading("Emerson Ruan Dantas Silva", level=1)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

contact = doc.add_paragraph()
contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
contact.add_run("Engenheiro de Software").bold = True
contact.add_run("\n")
contact.add_run("João Pessoa, PB  |  (83) 9 8685-3405  |  me@emersondantas.eti.br")
contact.add_run("\n")
contact.add_run("linkedin.com/in/dantasemerson  |  github.com/emersondantas")

# ==========================================
# PERFIL PROFISSIONAL
# ==========================================
add_separator()
doc.add_heading("Perfil Profissional", level=2)

doc.add_paragraph(
    "Engenheiro de Software com mais de 5 anos de experiência em desenvolvimento Full Stack, "
    "atuando com tecnologias como Java, Spring Boot, Angular, React, TypeScript, Docker, Kubernetes, "
    "Linux e Cloud Computing (AWS, GCP). Possuo MBA em Arquitetura de Software e Soluções, que "
    "fortaleceu minha visão sobre design de sistemas escaláveis e tomada de decisões técnicas. "
    "Atualmente cursando Mestrado em Engenharia de Software com foco em Dívidas Técnicas. "
    "Experiência em liderança técnica, revisão de código, mentoria de desenvolvedores e participação "
    "em decisões de arquitetura."
)

# ==========================================
# EXPERIÊNCIA PROFISSIONAL
# ==========================================
add_separator()
doc.add_heading("Experiência Profissional", level=2)

add_experience(
    "Engenheiro de Software Sênior", "Asaas", "Outubro de 2025 - Abril de 2026",
    context="Atuação no time de KYC (Know Your Customer), responsável pela obtenção, tratamento e manutenção das informações dos clientes, além de fluxos de encerramento de contas.",
    items=[
        "Desenvolvimento e manutenção de aplicações web utilizando Grails (Groovy) e GSP, integrando Design Systems baseados em React via Webpack",
        "Manutenção e desenvolvimento de funcionalidades no monolito e serviços satélites Grails, com foco em segurança e escalabilidade",
        "Manutenção em scripts SQL para MySQL, Jobs de processamento e cache com Redis",
        "Manutenção de integração com serviços terceiros para OCR",
        "Revisões de código de PRs do time e dos times da mesma tribo",
        "Participação em decisões técnicas de arquitetura e desenvolvimento",
        "Manutenção de fluxos assíncronos de jobs e webhooks para análises de contas",
        "Conformidade com práticas de desenvolvimento seguro em todos os ciclos de vida de software",
    ],
    techs=["Grails", "Groovy", "React", "Webpack", "MySQL", "Redis", "GSP", "KYC"],
)

add_experience(
    "Engenheiro de Software Pleno", "Strada", "Agosto de 2022 - Setembro de 2025",
    context="Atuação em múltiplos projetos do ecossistema logístico e financeiro: controle de fretes de grandes empresas agrícolas (documentação regulatória ao pagamento), gestão de frotas de caminhões (combustíveis, motoristas e rotas), leilões de cargas e conta digital (Pix, backend, backoffice e internet banking).",
    items=[
        "Desenvolvimento e manutenção de soluções usando tecnologias web e cloud (React, Angular, Spring Boot, Java, TypeScript, Oracle DB, Postgres, AWS, Cognito)",
        "Integração de sistemas com soluções de BaaS (banco como serviço)",
        "Resolução de chamados e incidentes de produção (War Rooms, N3)",
        "Manutenção em sistemas Java legados",
        "Testes unitários de API usando JUnit, Hamcrest e MockMVC",
        "Criação e manutenção de ambientes de testes, homologação e produção (Docker, Kubernetes, AWS, Cognito, S3)",
    ],
    techs=["React", "Angular", "Spring Boot", "Java", "TypeScript", "AWS", "Docker", "Kubernetes", "Oracle DB", "PostgreSQL", "Pix", "BaaS"],
)

add_experience(
    "Líder Técnico", "Phoebus Tecnologia", "Março de 2022 - Agosto de 2022",
    context="Atuação no desenvolvimento de um e-commerce integrado aos sistemas internos de fidelização de clientes da empresa.",
    items=[
        "Referência técnica no uso das tecnologias da equipe (Angular, Spring Boot, microsserviços, Java, TypeScript, PostgreSQL, MongoDB, Redis, Docker)",
        "Acompanhamento de desenvolvedores em suas atividades (dificuldades, revisões, sugestões)",
        "Refinamento de escopo de requisitos com tarefas técnicas",
        "Implementação de testes automatizados para APIs REST",
        "Proposta e aplicação de novas tecnologias, participando de decisões técnicas",
    ],
    techs=["Angular", "Spring Boot", "Java", "TypeScript", "PostgreSQL", "MongoDB", "Redis", "Docker", "E-commerce"],
)

add_experience(
    "Engenheiro de Software Pleno", "Phoebus Tecnologia", "Agosto de 2021 - Março de 2022",
    context="Atuação no desenvolvimento de um e-commerce integrado aos sistemas internos de fidelização de clientes da empresa.",
    items=[
        "Arquitetura e desenvolvimento de microsserviços usando Spring Boot com Java e a stack Spring Cloud",
        "Desenvolvimento de frontends com Angular e microfrontends com JsChannel (migração de Angular JS)",
        "Revisão de solicitações de merge e auxílio a outros desenvolvedores",
    ],
    techs=["Spring Boot", "Spring Cloud", "Angular", "Java", "Microsserviços", "E-commerce"],
)

add_experience(
    "Engenheiro de Software Júnior", "Phoebus Tecnologia", "Janeiro de 2021 - Agosto de 2021",
    items=[
        "Desenvolvimento de APIs usando Spring Boot com Java, prezando pelo desempenho e qualidade",
        "Scripts SQL para migração, versionamento e manipulação de banco de dados",
        "Criação e manutenção de frontends em Angular",
    ],
    techs=["Spring Boot", "Java", "Angular", "SQL"],
)

add_experience(
    "Desenvolvedor Backend Java - Trainee Pleno", "Phoebus Tecnologia", "Agosto a Dezembro de 2020",
    items=[
        "Desenvolvimento de APIs com Spring Boot, scripts SQL e frontends em Angular",
        "Cobertura de testes unitários com JUnit, Hamcrest e MockMVC",
    ],
    techs=["Spring Boot", "Java", "Angular", "JUnit"],
)

add_experience(
    "Consultor / Desenvolvedor Full Stack", "Projeto Tracy Technical Debt - UFPB", "Dezembro de 2019 - Agosto de 2021",
    items=[
        "Desenvolvimento e melhoria de funcionalidades para a ferramenta de gerência de dívidas técnicas Tracy TD",
        "Criação de ambientes de produção e testes em diferentes provedores cloud (On-Premise, AWS e Digital Ocean), usando Docker e Docker Compose",
    ],
    techs=["Angular", "Spring Boot", "PostgreSQL", "Docker", "AWS", "Digital Ocean"],
)

add_experience(
    "Diretor de Projetos", "Empresa Júnior - Eagle Soluções Digitais (UFPB)", "Março a Dezembro de 2019",
    items=[
        "Gerenciamento de projetos: acompanhamento de equipes, gerência de prazos, documentações, riscos e oportunidades",
    ],
)

add_experience(
    "Desenvolvedor Full Stack", "Empresa Júnior - Eagle Soluções Digitais (UFPB)", "Dezembro de 2018 - Março de 2019",
    items=[
        "Desenvolvimento de soluções computacionais com Android Nativo, Django, Spring Boot, PostgreSQL",
        "Criação de ambientes de produção usando servidores Linux e Nginx",
    ],
    techs=["Android", "Django", "Spring Boot", "PostgreSQL", "AWS EC2", "Nginx"],
)

add_experience(
    "Desenvolvedor / Pesquisador", "Projeto Apps4Society - UFPB", "Dezembro de 2018 - Dezembro de 2019",
    items=[
        "Prototipação e desenvolvimento de aplicativos Android (Java e Kotlin)",
        "Desenvolvimento de API em Spring Boot",
    ],
    techs=["Android", "Java", "Kotlin", "Spring Boot"],
)

add_experience(
    "Tutor Java", "UFPB - Campus IV - DCX", "Março a Novembro de 2019",
    items=["Orientação e ensino de Linguagem de Programação com Java"],
)

add_experience(
    "Monitor Python", "UFPB - Campus IV - DCX", "Março a Novembro de 2018",
    items=["Orientação e ensino de Introdução à Programação com Python"],
)

# ==========================================
# PROJETOS
# ==========================================
add_separator()
doc.add_heading("Projetos", level=2)

add_project(
    "Aurum.rocks", "Criador e Mantenedor", "Projeto Pessoal",
    "Gestor de finanças pessoais focado na visão a médio e longo prazo e no controle de gastos. Oferece timeline de transações com filtros, resumo financeiro, balanço diário, gestão de reservas, alertas de saldos negativos e gráficos de relatório mensal.",
    items=[
        "Frontend em React com Next.js e design Neo Brutalism",
        "Backend em Spring Boot com PostgreSQL e cache Caffeine",
        "Integração com LLMs para insights financeiros",
        "Autenticação e autorização via AWS Cognito com SSO Google",
    ],
    techs=["React", "Next.js", "Spring Boot", "PostgreSQL", "Caffeine", "LLMs", "AWS Cognito", "Neo Brutalism"],
)

add_project(
    "Seustand.com.br", "Arquiteto de Software e Líder Técnico", "Agosto de 2023 - Março de 2024",
    "Plataforma que conecta organizadores de eventos a expositores. Permite a visualização de eventos, reserva e compra de stands para exposição, com fluxo completo de pagamentos integrado — incluindo parcelamento, múltiplas formas de pagamento e sistema de devoluções. Os organizadores recebem os valores diretamente pela plataforma.",
    items=[
        "Construção do zero: arquitetura, desenvolvimento e liderança técnica da equipe",
        "Desenvolvimento do frontend (Angular 12), backend (API Spring Boot 3.x) e banco de dados (PostgreSQL)",
        "Integração com gateway de pagamentos com suporte a parcelamento, múltiplos métodos e devoluções",
        "Autenticação e autorização via AWS Cognito",
        "Pipelines de CI/CD (GitHub Actions) e infraestrutura AWS (EC2, Docker, Route53, S3)",
    ],
    techs=["Angular", "Spring Boot", "PostgreSQL", "AWS", "AWS Cognito", "Docker", "GitHub Actions"],
)

add_project(
    "InovaturPB - Observatório de Turismo da Paraíba", "Arquiteto e Engenheiro de Software", "Março de 2023 - Agosto de 2023",
    "Sistema de gestão de atrativos turísticos do estado da Paraíba. Facilita a busca e a visualização de informações sobre pontos turísticos, além de oferecer um backoffice para a gestão centralizada dos dados pelo órgão responsável.",
    items=[
        "Construção do zero: da arquitetura ao desenvolvimento completo",
        "Frontend (React), backend (API Spring Boot 3.x) e banco de dados (PostgreSQL)",
        "Estratégias de microfrontend para integração com portal WordPress existente",
        "Implementação de cache com Redis para integração com serviços terceiros",
    ],
    techs=["React", "Spring Boot", "PostgreSQL", "Redis", "Microfrontend", "WordPress"],
)

add_project(
    "Shaders Design", "Arquiteto de Software e Líder Técnico", "Dezembro de 2021 - Julho de 2022",
    "Marketplace de assets 3D (malhas) voltado para profissionais de design e arquitetura. Os usuários podem navegar, visualizar e adquirir malhas para uso em seus projetos, com pagamento integrado à plataforma.",
    items=[
        "Construção do zero: arquitetura, desenvolvimento e liderança técnica da equipe",
        "Desenvolvimento da API em Spring Boot + PostgreSQL e frontend em Angular",
        "Integração com gateway de pagamentos para compra de assets",
        "CI/CD de homologação (GitHub Actions) e produção na AWS (Elastic Beanstalk, API Gateway, Load Balancer, S3, RDS)",
    ],
    techs=["Angular", "Spring Boot", "PostgreSQL", "AWS", "GitHub Actions", "Elastic Beanstalk"],
)

add_project(
    "MyFastSurvey.com.br", "Arquiteto de Software e Líder de Equipe", "Fevereiro de 2020 - Presente",
    "Plataforma completa de pesquisas presenciais e online. Permite criar questionários com diversos tipos de perguntas, escalar entrevistadores que utilizam o aplicativo para realizar pesquisas em campo e eventos, com auditoria completa — incluindo gravação de áudio de cada entrevista. Também oferece questionários online com links únicos e envio em massa por e-mail, relatórios detalhados e exportação de dados em XLSX, PDF e DOCX.",
    items=[
        "Construção do zero: arquitetura, desenvolvimento e liderança técnica da equipe. Manutenção e melhorias contínuas até hoje",
        "Arquitetura da API (Spring Boot) e frontend (Angular)",
        "Sistema de auditoria com gravação de áudio das entrevistas",
        "Relatórios e exportação de dados em múltiplos formatos (XLSX, PDF, DOCX)",
        "Integração com sistema de pagamentos para compra de créditos de entrevistas",
        "Definição de processos de CI/CD e infraestrutura em DigitalOcean com Google Cloud Storage",
    ],
    techs=["Angular", "Spring Boot", "Docker", "DigitalOcean", "Google Cloud Storage"],
)

# ==========================================
# FORMAÇÃO ACADÊMICA
# ==========================================
add_separator()
doc.add_heading("Formação Acadêmica", level=2)

edu_items = [
    ("Mestrado em Informática - Engenharia de Software", "Universidade Tecnológica Federal do Paraná (UTFPR)", "Fevereiro de 2024 - Presente", "Foco em Dívidas Técnicas"),
    ("MBA em Arquitetura de Software e Soluções", "XP Educação", "Agosto de 2022 - Julho de 2023", "Pós-graduação Lato Sensu"),
    ("Bacharel em Sistemas de Informação", "Universidade Federal da Paraíba (UFPB)", "Julho de 2017 - Julho de 2021", "CRA: 9.1"),
]

for degree, school, period, detail in edu_items:
    h = doc.add_heading(degree, level=3)
    h.paragraph_format.space_before = Pt(8)
    p = doc.add_paragraph()
    run_school = p.add_run(school)
    run_school.bold = True
    run_school.font.size = Pt(10.5)
    p.add_run(f"  |  {period}").font.size = Pt(10)
    if detail:
        doc.add_paragraph(detail).paragraph_format.space_before = Pt(0)

# ==========================================
# HABILIDADES TÉCNICAS
# ==========================================
add_separator()
doc.add_heading("Habilidades Técnicas", level=2)

skills_sections = [
    ("Backend", "Java, Spring Boot, Groovy, Grails, Python, Django, Spring Security, Spring Cloud Gateway, Spring Cloud Eureka, OpenFeign, Flyway, JUnit, Hamcrest, MockMVC"),
    ("Frontend", "Angular, React, Next.js, TypeScript, JavaScript, HTML5, CSS3, Angular Material, Chakra UI, Ant Design, Neo Brutalism, Bootstrap, Sass"),
    ("Banco de Dados e Cache", "PostgreSQL, MySQL, Oracle, MongoDB, Redis, Caffeine"),
    ("DevOps e Cloud", "Docker, Docker Compose, Kubernetes, Linux, Nginx, GitHub Actions, GitLab CI, TeamCity, Grafana, Prometheus, Portainer, SonarQube, Servidores Linux"),
    ("AWS", "Cognito, S3, EC2, Route 53, API Gateway, NLB, ALB, Elastic Beanstalk, RDS"),
    ("Google Cloud", "Cloud Storage, Maps API"),
    ("Inteligência Artificial", "Gemini, ChatGPT, Claude Code, Llama (modelos locais), Gemini API, OpenAI API, Gestão de tokens, Engenharia de prompts, Skills e agentes"),
    ("Ferramentas", "Git, GitHub, GitLab, IntelliJ, VS Code, Gradle, DBeaver, Amazon Q, GitHub Copilot, Kiro"),
]

for category, skills in skills_sections:
    p = doc.add_paragraph()
    run_cat = p.add_run(f"{category}: ")
    run_cat.bold = True
    run_cat.font.size = Pt(10.5)
    p.add_run(skills).font.size = Pt(10.5)

# ==========================================
# SOFT SKILLS
# ==========================================
add_separator()
doc.add_heading("Soft Skills", level=2)

doc.add_paragraph(
    "Liderança técnica, Code review, Feedback construtivo, Mentoria de desenvolvedores, "
    "Tomada de decisão técnica, Comunicação clara, Planejamento e estimativas, "
    "Resolução de problemas, Trabalho em equipe, Autonomia, Gestão de prioridades, "
    "Documentação técnica"
)

# ==========================================
# REFERÊNCIAS
# ==========================================
add_separator()
doc.add_heading("Referências", level=2)

ref1 = doc.add_paragraph()
ref1.add_run("Patrick Onofre").bold = True
ref1.add_run("\nEngenheiro de Software Especialista - Strada")
ref1.add_run("\nonofre.patrick@gmail.com  |  linkedin.com/in/patrick-onofre-2b658a34")

ref2 = doc.add_paragraph()
ref2.add_run("Rodrigo Rebouças").bold = True
ref2.add_run("\nProfessor, PhD - Universidade Federal da Paraíba (UFPB)")
ref2.add_run("\nrodrigor@dcx.ufpb.br  |  linkedin.com/in/rodrigoreboucas")

# ==========================================
# SALVAR
# ==========================================
output_path = "download/curriculo_emerson_ruan_dantas_silva.docx"
doc.save(output_path)
print(f"DOCX gerado com sucesso: {output_path}")
